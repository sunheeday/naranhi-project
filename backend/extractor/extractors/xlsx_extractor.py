from __future__ import annotations

from pathlib import Path

import openpyxl

from extractor.models import ExtractedText


MAX_ROWS_PER_SHEET = 5000
MAX_COLS_PER_ROW = 200


async def extract_xlsx_text(
    path: Path,
    *,
    source_name: str,
) -> ExtractedText:
    """xlsx 파일의 셀 값을 TSV 형식 텍스트로 추출.

    - 보이는 시트만 처리 (숨김 시트 스킵)
    - 시트마다 [Sheet: 이름] 헤더 출력
    - 각 행: 셀 값을 탭(\t)으로 구분
    - 빈 행은 스킵
    - 시트당 최대 MAX_ROWS_PER_SHEET행, 행당 최대 MAX_COLS_PER_ROW열
    """
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises varied error types.
        return ExtractedText(
            source=source_name,
            method="xlsx_openpyxl",
            text="",
            status="extract_failed",
            warnings=[f"xlsx_load_failed: {type(exc).__name__}: {exc}"],
        )

    parts: list[str] = []
    truncated_sheets: list[str] = []

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if ws.sheet_state != "visible":
                continue

            sheet_lines: list[str] = []
            row_count = 0
            truncated = False
            for row in ws.iter_rows(values_only=True):
                if row_count >= MAX_ROWS_PER_SHEET:
                    truncated = True
                    break
                cells = [_cell_to_str(c) for c in row[:MAX_COLS_PER_ROW]]
                if not any(cell.strip() for cell in cells):
                    continue
                sheet_lines.append("\t".join(cells))
                row_count += 1

            if sheet_lines:
                parts.append(f"[Sheet: {sheet_name}]")
                parts.extend(sheet_lines)
            if truncated:
                truncated_sheets.append(sheet_name)
    finally:
        wb.close()

    text = "\n".join(parts)
    warnings: list[str] = []
    if truncated_sheets:
        warnings.append(
            f"truncated_sheets={','.join(truncated_sheets)} "
            f"(>{MAX_ROWS_PER_SHEET} rows)"
        )

    return ExtractedText(
        source=source_name,
        method="xlsx_openpyxl",
        text=text,
        status="success" if text.strip() else "empty_or_unreadable",
        warnings=warnings,
    )


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    return str(value)
